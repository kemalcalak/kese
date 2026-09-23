"""Unit tests for report and export calculations."""

import asyncio
from collections.abc import AsyncIterator
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from typing import cast
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from kese.api.exports import csv_chunks, stream_csv_export
from kese.services.exports import build_xlsx
from kese.services.reports import export_range, parse_month


def test_month_and_export_boundaries_are_utc() -> None:
    assert parse_month("2026-05").isoformat() == "2026-05-01"
    assert export_range("2026-05-01", "2026-05-31") == (
        datetime(2026, 5, 1, tzinfo=UTC),
        datetime(2026, 6, 1, tzinfo=UTC),
    )


def test_csv_chunks_have_decimal_text_and_expected_header() -> None:
    async def rows() -> AsyncIterator[tuple[datetime, str, str, str | None, Decimal]]:
        yield (
            datetime(2026, 5, 2, 23, 0, tzinfo=UTC),
            "Vadesiz",
            "MIGROS",
            "Market",
            Decimal("-450.75"),
        )

    async def collect() -> list[str]:
        return [chunk async for chunk in csv_chunks(rows())]

    assert "".join(asyncio.run(collect())) == (
        "date,account,description,category,amount\r\n"
        "2026-05-02,Vadesiz,MIGROS,Market,-450.75\r\n"
    )


def test_csv_export_keeps_session_open_until_stream_finishes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class SessionLifecycle:
        open = False
        closed = False

        def __call__(self) -> SessionLifecycle:
            return self

        async def __aenter__(self) -> AsyncSession:
            self.open = True
            return cast(AsyncSession, self)

        async def __aexit__(
            self,
            exc_type: object,
            exc_value: object,
            traceback: object,
        ) -> None:
            self.closed = True

    lifecycle = SessionLifecycle()

    async def fake_stream(
        session: AsyncSession, user_id: UUID, since: str, until: str
    ) -> AsyncIterator[tuple[datetime, str, str, str | None, Decimal]]:
        assert session is cast(AsyncSession, lifecycle)
        assert lifecycle.open is True
        assert lifecycle.closed is False
        yield (
            datetime(2026, 5, 2, tzinfo=UTC),
            "Vadesiz",
            "MIGROS",
            None,
            Decimal(1),
        )

    monkeypatch.setattr("kese.api.exports.stream_export_rows", fake_stream)

    async def collect() -> list[str]:
        factory = cast(async_sessionmaker[AsyncSession], lifecycle)
        return [
            chunk
            async for chunk in stream_csv_export(
                factory, uuid4(), "2026-05-01", "2026-05-02"
            )
        ]

    assert "".join(asyncio.run(collect())).endswith("MIGROS,,1\r\n")
    assert lifecycle.closed is True


def test_xlsx_export_writes_amount_as_a_numeric_cell() -> None:
    rows = [
        (
            datetime(2026, 5, 2, 9, 0, tzinfo=UTC),
            "Vadesiz",
            "MIGROS",
            None,
            Decimal("-450.75"),
        )
    ]
    sheet = load_workbook(BytesIO(build_xlsx(rows))).active
    assert sheet is not None
    assert sheet.cell(row=1, column=5).value == "amount"
    assert sheet.cell(row=2, column=5).value == Decimal("-450.75")
