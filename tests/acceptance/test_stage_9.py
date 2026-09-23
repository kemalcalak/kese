"""Stage 9 - reports and exports.

Written by hand before the stage and never edited: it states what the stage
has to do. It needs the compose database up and the migrations applied.

Every test uses its own user, and reports only ever read the caller's rows,
so no test can see another's data - the lesson of stage 7's shared cache.
"""

from __future__ import annotations

import csv
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Any, cast
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from kese.main import create_app

PASSWORD = "correct horse battery staple"

Headers = dict[str, str]


def an_email() -> str:
    return f"user-{uuid4().hex}@example.com"


def a_user(client: TestClient) -> Headers:
    email = an_email()
    client.post("/auth/register", json={"email": email, "password": PASSWORD})
    tokens = client.post(
        "/auth/login", json={"email": email, "password": PASSWORD}
    ).json()
    return {"Authorization": f"Bearer {tokens['access_token']}"}


def an_account(
    client: TestClient, headers: Headers, name: str = "Vadesiz", currency: str = "TRY"
) -> str:
    response = client.post(
        "/accounts", json={"name": name, "currency": currency}, headers=headers
    )
    assert response.status_code == 201, response.text
    return cast(str, response.json()["id"])


def a_category(client: TestClient, headers: Headers, name: str) -> str:
    response = client.post("/categories", json={"name": name}, headers=headers)
    assert response.status_code == 201, response.text
    return cast(str, response.json()["id"])


def book(
    client: TestClient,
    headers: Headers,
    account_id: str,
    amount: str,
    on: str,
    category_id: str | None = None,
    description: str = "hareket",
) -> None:
    body: dict[str, Any] = {
        "amount": amount,
        "occurred_at": on,
        "description": description,
    }
    if category_id is not None:
        body["category_id"] = category_id
    response = client.post(
        f"/accounts/{account_id}/transactions", json=body, headers=headers
    )
    assert response.status_code == 201, response.text


def monthly(
    client: TestClient, headers: Headers, month: str = "2026-05", currency: str = "TRY"
) -> dict[str, Any]:
    response = client.get(
        "/reports/monthly",
        params={"month": month, "currency": currency},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return cast(dict[str, Any], response.json())


def dec(value: Any) -> Decimal:
    return Decimal(str(value))


def test_the_month_adds_up() -> None:
    with TestClient(create_app()) as client:
        headers = a_user(client)
        account = an_account(client, headers)
        book(client, headers, account, "32000.00", "2026-05-01T09:00:00Z")
        book(client, headers, account, "-450.75", "2026-05-02T09:00:00Z")
        book(client, headers, account, "-1200.00", "2026-05-15T09:00:00Z")
        book(client, headers, account, "-999.00", "2026-06-01T00:00:00Z")

        report = monthly(client, headers)

    assert report["month"] == "2026-05"
    assert dec(report["income"]) == Decimal("32000.00")
    assert dec(report["expenses"]) == Decimal("1650.75")
    assert dec(report["net"]) == Decimal("30349.25")


def test_categories_are_ranked_by_what_was_spent() -> None:
    with TestClient(create_app()) as client:
        headers = a_user(client)
        account = an_account(client, headers)
        rent = a_category(client, headers, "Kira")
        market = a_category(client, headers, "Market")
        bills = a_category(client, headers, "Fatura")
        book(client, headers, account, "-500.00", "2026-05-01T09:00:00Z", rent)
        book(client, headers, account, "-300.00", "2026-05-02T09:00:00Z", market)
        book(client, headers, account, "-300.00", "2026-05-03T09:00:00Z", bills)

        rows = monthly(client, headers)["by_category"]

    assert [(row["name"], row["rank"]) for row in rows] == [
        ("Kira", 1),
        ("Fatura", 2),
        ("Market", 2),
    ]
    assert [dec(row["spent"]) for row in rows] == [
        Decimal("500.00"),
        Decimal("300.00"),
        Decimal("300.00"),
    ]


def test_each_category_carries_its_share_of_the_spending() -> None:
    with TestClient(create_app()) as client:
        headers = a_user(client)
        account = an_account(client, headers)
        rent = a_category(client, headers, "Kira")
        market = a_category(client, headers, "Market")
        book(client, headers, account, "-700.00", "2026-05-01T09:00:00Z", rent)
        book(client, headers, account, "-300.00", "2026-05-02T09:00:00Z", market)

        rows = monthly(client, headers)["by_category"]

    shares = {row["name"]: dec(row["share"]) for row in rows}
    assert shares == {"Kira": Decimal("0.7"), "Market": Decimal("0.3")}


def test_uncategorised_spending_is_its_own_row() -> None:
    with TestClient(create_app()) as client:
        headers = a_user(client)
        account = an_account(client, headers)
        rent = a_category(client, headers, "Kira")
        book(client, headers, account, "-500.00", "2026-05-01T09:00:00Z", rent)
        book(client, headers, account, "-80.00", "2026-05-02T09:00:00Z")

        rows = monthly(client, headers)["by_category"]

    uncategorised = [row for row in rows if row["category_id"] is None]
    assert len(uncategorised) == 1
    assert dec(uncategorised[0]["spent"]) == Decimal("80.00")


def test_a_report_reads_only_the_callers_money() -> None:
    with TestClient(create_app()) as client:
        stranger = a_user(client)
        book(
            client,
            stranger,
            an_account(client, stranger),
            "-5000.00",
            "2026-05-01T09:00:00Z",
        )

        headers = a_user(client)
        book(
            client,
            headers,
            an_account(client, headers),
            "-10.00",
            "2026-05-01T09:00:00Z",
        )

        report = monthly(client, headers)

    assert dec(report["expenses"]) == Decimal("10.00")


def test_a_report_stays_in_one_currency() -> None:
    with TestClient(create_app()) as client:
        headers = a_user(client)
        lira = an_account(client, headers, "Vadesiz", "TRY")
        dollars = an_account(client, headers, "Döviz", "USD")
        book(client, headers, lira, "-100.00", "2026-05-01T09:00:00Z")
        book(client, headers, dollars, "-40.00", "2026-05-01T09:00:00Z")

        in_lira = monthly(client, headers, currency="TRY")
        in_dollars = monthly(client, headers, currency="USD")

    assert dec(in_lira["expenses"]) == Decimal("100.00")
    assert dec(in_dollars["expenses"]) == Decimal("40.00")


def test_the_trend_keeps_empty_months_and_a_running_total() -> None:
    with TestClient(create_app()) as client:
        headers = a_user(client)
        account = an_account(client, headers)
        book(client, headers, account, "1000.00", "2026-03-10T09:00:00Z")
        book(client, headers, account, "-400.00", "2026-05-10T09:00:00Z")

        response = client.get(
            "/reports/trend",
            params={"until": "2026-05", "months": 3, "currency": "TRY"},
            headers=headers,
        )

    assert response.status_code == 200, response.text
    series = response.json()
    assert [point["month"] for point in series] == ["2026-03", "2026-04", "2026-05"]
    assert [dec(point["net"]) for point in series] == [
        Decimal("1000.00"),
        Decimal(0),
        Decimal("-400.00"),
    ]
    assert [dec(point["running_net"]) for point in series] == [
        Decimal("1000.00"),
        Decimal("1000.00"),
        Decimal("600.00"),
    ]


def test_the_csv_export_is_the_months_rows() -> None:
    with TestClient(create_app()) as client:
        headers = a_user(client)
        account = an_account(client, headers, "Vadesiz")
        market = a_category(client, headers, "Market")
        book(
            client,
            headers,
            account,
            "-450.75",
            "2026-05-02T09:00:00Z",
            market,
            "MIGROS",
        )
        book(client, headers, account, "-10.00", "2026-04-30T09:00:00Z")

        response = client.get(
            "/exports/transactions.csv",
            params={"since": "2026-05-01", "until": "2026-05-31"},
            headers=headers,
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    rows = list(csv.DictReader(StringIO(response.text)))
    assert list(rows[0].keys()) == [
        "date",
        "account",
        "description",
        "category",
        "amount",
    ]
    assert len(rows) == 1
    assert rows[0]["date"] == "2026-05-02"
    assert rows[0]["account"] == "Vadesiz"
    assert rows[0]["category"] == "Market"
    assert Decimal(rows[0]["amount"]) == Decimal("-450.75")


def test_the_xlsx_export_opens_as_a_workbook() -> None:
    with TestClient(create_app()) as client:
        headers = a_user(client)
        account = an_account(client, headers)
        book(client, headers, account, "-450.75", "2026-05-02T09:00:00Z")
        book(client, headers, account, "32000.00", "2026-05-01T09:00:00Z")

        response = client.get(
            "/exports/transactions.xlsx",
            params={"since": "2026-05-01", "until": "2026-05-31"},
            headers=headers,
        )

    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet is not None
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == ["date", "account", "description", "category", "amount"]
    assert len(rows) == 3
    assert {Decimal(str(row[4])) for row in rows[1:]} == {
        Decimal("-450.75"),
        Decimal("32000.00"),
    }


def test_an_export_carries_only_the_callers_rows() -> None:
    with TestClient(create_app()) as client:
        stranger = a_user(client)
        book(
            client,
            stranger,
            an_account(client, stranger),
            "-5000.00",
            "2026-05-01T09:00:00Z",
        )

        headers = a_user(client)
        response = client.get(
            "/exports/transactions.csv",
            params={"since": "2026-05-01", "until": "2026-05-31"},
            headers=headers,
        )

    assert list(csv.DictReader(StringIO(response.text))) == []


def test_reports_and_exports_need_a_token() -> None:
    with TestClient(create_app()) as client:
        report = client.get("/reports/monthly", params={"month": "2026-05"})
        export = client.get("/exports/transactions.csv")

    assert report.status_code == 401
    assert export.status_code == 401
