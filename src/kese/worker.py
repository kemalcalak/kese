"""Worker for processing queued statement imports."""

from __future__ import annotations

import csv
import hashlib
from collections.abc import Iterable, Iterator
from datetime import UTC, date, datetime, time
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Any, cast
from zipfile import BadZipFile

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from kese.core.database import create_engine
from kese.core.settings import Settings
from kese.models import Statement
from kese.services.transactions import import_transaction

Row = tuple[datetime, str, Decimal]
ParsedRow = Row | None


def _aware_datetime(value: Any) -> datetime:
    """Parse a supported date value and normalize it to UTC."""
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    elif isinstance(value, str):
        parsed = datetime.fromisoformat(value.strip())
    else:
        raise TypeError("invalid date")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed.astimezone(UTC)


def _row(values: dict[str, Any]) -> Row:
    """Convert one mapped statement row to typed import values."""
    raw_date = values.get("date")
    raw_description = values.get("description")
    raw_amount = values.get("amount")
    if raw_description is None or not str(raw_description).strip():
        raise ValueError("invalid description")
    if raw_amount is None:
        raise ValueError("invalid amount")
    return (
        _aware_datetime(raw_date),
        str(raw_description),
        Decimal(str(raw_amount).strip()),
    )


def _parse_csv(content: bytes) -> Iterator[ParsedRow]:
    """Yield typed rows from a CSV statement, tolerating malformed rows."""
    reader = csv.DictReader(StringIO(content.decode("utf-8-sig"), newline=""))
    if reader.fieldnames is None or set(reader.fieldnames) != {
        "date",
        "description",
        "amount",
    }:
        raise ValueError("invalid statement header")
    for values in reader:
        try:
            yield _row(values)
        except TypeError, ValueError, ArithmeticError:
            yield None


def _parse_xlsx(content: bytes) -> Iterator[ParsedRow]:
    """Yield typed rows from an XLSX statement."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        if worksheet is None:
            raise ValueError("missing worksheet")
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None or tuple(header) != ("date", "description", "amount"):
            raise ValueError("invalid statement header")
        for values in rows:
            if all(value is None for value in values):
                continue
            if len(values) != 3:
                yield None
                continue
            try:
                yield _row(dict(zip(("date", "description", "amount"), values)))
            except TypeError, ValueError, ArithmeticError:
                yield None
    finally:
        workbook.close()


def parse_statement(filename: str, content: bytes) -> Iterable[Row]:
    """Parse a supported statement format into typed rows."""
    if filename.lower().endswith(".csv"):
        return cast(Iterable[Row], _parse_csv(content))
    if filename.lower().endswith(".xlsx"):
        return cast(Iterable[Row], _parse_xlsx(content))
    raise ValueError("unsupported statement format")


def _fingerprint(
    account_id: Any, occurred_at: datetime, description: str, amount: Decimal
) -> str:
    """Build the stable database-backed duplicate key for one row."""
    value = f"{account_id}|{occurred_at.isoformat()}|{description}|{amount}"
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _is_fingerprint_conflict(error: IntegrityError) -> bool:
    """Identify the transaction fingerprint unique constraint violation."""
    return "fingerprint" in str(error.orig).casefold()


async def _process_statement(session: AsyncSession, statement: Statement) -> None:
    """Import all rows for a locked statement job."""
    imported = duplicates = failed = 0
    try:
        rows = parse_statement(statement.filename, statement.content)
        for raw_row in rows:
            if raw_row is None:
                failed += 1
                continue
            try:
                occurred_at, description, amount = raw_row
                async with session.begin_nested():
                    transaction = await import_transaction(
                        session,
                        statement.user_id,
                        statement.account_id,
                        amount,
                        occurred_at,
                        description,
                        _fingerprint(
                            statement.account_id, occurred_at, description, amount
                        ),
                    )
                    if transaction is None:
                        raise ValueError("statement account is not owned by its user")
                imported += 1
            except IntegrityError as error:
                if not _is_fingerprint_conflict(error):
                    raise
                duplicates += 1
            except TypeError, ValueError, ArithmeticError:
                failed += 1
    except BadZipFile, TypeError, ValueError, ArithmeticError:
        failed += 1
        statement.status = "failed"
    else:
        statement.status = "done"
    statement.imported = imported
    statement.duplicates = duplicates
    statement.failed = failed


async def run_once() -> int:
    """Claim and process at most one pending statement, returning job count."""
    engine = create_engine(Settings().database_url)
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        async with session_factory() as session:
            async with session.begin():
                result = await session.execute(
                    select(Statement)
                    .where(Statement.status == "pending")
                    .order_by(Statement.created_at.desc(), Statement.id.desc())
                    .with_for_update(skip_locked=True)
                    .limit(1)
                )
                statement = result.scalar_one_or_none()
                if statement is None:
                    return 0
                await _process_statement(session, statement)
            return 1
    finally:
        await engine.dispose()
